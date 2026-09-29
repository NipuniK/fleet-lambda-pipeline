import docx
import re

doc = docx.Document()
with open('EC8203_Fleet_Ops_Report.md', 'r') as f:
    lines = f.readlines()

for line in lines:
    line = line.strip()
    if not line:
        continue
    if line.startswith('#'):
        level = len(line) - len(line.lstrip('#'))
        text = line.lstrip('#').strip()
        doc.add_heading(text, level=level)
    elif line.startswith('```'):
        continue
    elif line.startswith('Figure'):
        p = doc.add_paragraph(line)
        p.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
    else:
        doc.add_paragraph(line)

doc.save('EC8203_Fleet_Ops_Report_Final.docx')
print("DOCX created successfully.")
